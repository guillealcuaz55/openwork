from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\matrix")
DOCX_PATH = ROOT / "Brain - Producto y Arquitectura Base.docx"
OUT_DIR = ROOT / "docx_render_brain"
DIAGRAM_PATH = OUT_DIR / "brain_architecture_diagram.png"


def font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline, radius=18, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, box, text, fnt, fill="#172026", line_gap=6):
    x1, y1, x2, y2 = box
    lines = []
    for raw in text.split("\n"):
        words = raw.split(" ")
        current = ""
        for word in words:
            trial = word if not current else f"{current} {word}"
            if draw.textbbox((0, 0), trial, font=fnt)[2] <= (x2 - x1 - 32):
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2 - 2
    for line, h in zip(lines, heights):
        w = draw.textbbox((0, 0), line, font=fnt)[2]
        draw.text((x1 + ((x2 - x1) - w) / 2, y), line, font=fnt, fill=fill)
        y += h + line_gap


def arrow(draw, start, end, color="#56616b", width=3):
    draw.line([start, end], fill=color, width=width)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) > abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 13 * direction, ey - 7), (ex - 13 * direction, ey + 7)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 7, ey - 13 * direction), (ex + 7, ey - 13 * direction)]
    draw.polygon(points, fill=color)


def build_diagram():
    OUT_DIR.mkdir(exist_ok=True)
    scale = 2
    w, h = 1700, 1180
    img = Image.new("RGB", (w, h), "#f7f5f0")
    draw = ImageDraw.Draw(img)

    title_font = font(42, True)
    label_font = font(24, True)
    small_font = font(22, False)

    draw.text((70, 48), "Brain: arquitectura operativa con MCP", font=title_font, fill="#111820")
    draw.text(
        (72, 106),
        "El cliente trabaja desde su interfaz habitual; MCP conecta agentes, memoria y herramientas reales con revisión humana.",
        font=small_font,
        fill="#5a646f",
    )

    boxes = {
        "team": (90, 190, 395, 280, "Equipo cliente", "#ffffff"),
        "ui": (520, 180, 1120, 290, "Claude Desktop / OpenWork\nCursor / ChatGPT", "#ffffff"),
        "mcp": (690, 390, 950, 495, "MCP Server\nprivado", "#16202a"),
        "brain": (120, 610, 470, 715, "Company Brain", "#fff8df"),
        "tools": (680, 610, 980, 715, "Conectores", "#e9f4ff"),
        "agents": (1190, 610, 1535, 715, "Agentes\nespecializados", "#edf8ee"),
        "review": (1190, 830, 1535, 925, "Review queue", "#fff1e8"),
        "human": (1190, 990, 1535, 1080, "Humano\naprueba", "#ffffff"),
    }

    for _, (x1, y1, x2, y2, txt, fill) in boxes.items():
        outline = "#c9c1b2" if fill != "#16202a" else "#16202a"
        rounded_box(draw, (x1, y1, x2, y2), fill, outline, radius=22, width=2)
        centered_text(draw, (x1, y1, x2, y2), txt, label_font, "#ffffff" if fill == "#16202a" else "#172026")

    # Main arrows
    arrow(draw, (395, 235), (520, 235))
    arrow(draw, (820, 290), (820, 390))
    arrow(draw, (690, 448), (470, 650))
    arrow(draw, (820, 495), (820, 610))
    arrow(draw, (950, 448), (1190, 650))
    arrow(draw, (1362, 715), (1362, 830))
    arrow(draw, (1362, 925), (1362, 990))
    arrow(draw, (1190, 1035), (470, 680))

    # Brain children
    brain_items = [
        (95, 770, 500, 835, "knowledge/ procesos, políticas, criterios"),
        (95, 850, 500, 915, "skills/ procedimientos ejecutables"),
        (95, 930, 500, 995, "memory/ logs y aprendizajes"),
    ]
    for item in brain_items:
        rounded_box(draw, item[:4], "#ffffff", "#d9cfbc", radius=16, width=2)
        centered_text(draw, item[:4], item[4], small_font, "#172026")
        arrow(draw, (295, 715), (295, item[1]))

    # Tool children
    tool_items = [
        (610, 770, 760, 835, "Drive /\nSharePoint"),
        (780, 770, 930, 835, "Gmail /\nOutlook"),
        (950, 770, 1100, 835, "Calendar"),
        (610, 875, 835, 950, "CRM / ERP /\nHelpdesk / Sheets"),
        (860, 875, 1100, 950, "Slack / Teams /\nWhatsApp"),
    ]
    for item in tool_items:
        rounded_box(draw, item[:4], "#ffffff", "#c7d8e8", radius=16, width=2)
        centered_text(draw, item[:4], item[4], small_font, "#172026")
        arrow(draw, (830, 715), ((item[0] + item[2]) // 2, item[1]))

    # Legend
    rounded_box(draw, (1190, 190, 1535, 330), "#ffffff", "#d8d0c2", radius=18, width=2)
    centered_text(
        draw,
        (1210, 205, 1515, 315),
        "Regla base:\nleer, analizar y proponer;\nejecutar solo con permiso.",
        small_font,
        "#172026",
    )

    img = img.resize((w // scale, h // scale), Image.Resampling.LANCZOS)
    img.save(DIAGRAM_PATH)


def insert_into_docx():
    doc = Document(DOCX_PATH)
    start_idx = None
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "Arquitectura visual de Brain":
            start_idx = idx
            break
    if start_idx is not None:
        # Remove the prior inserted diagram section so the script is idempotent.
        body = doc._body._element
        elements_to_remove = []
        first_element = doc.paragraphs[max(start_idx - 1, 0)]._element
        found = False
        for child in list(body):
            if child is first_element:
                found = True
            if found and child.tag.endswith("}sectPr"):
                break
            if found:
                elements_to_remove.append(child)
        for element in elements_to_remove:
            body.remove(element)

    doc.add_page_break()
    heading = doc.add_heading("Arquitectura visual de Brain", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    intro = doc.add_paragraph()
    intro.add_run(
        "Este diagrama resume la capa operativa: el equipo trabaja desde sus interfaces habituales, "
        "el MCP server conecta el Brain con herramientas reales y los agentes producen trabajo revisable "
        "antes de cualquier accion sensible."
    )
    intro.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(DIAGRAM_PATH), width=Inches(6.9))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = caption.add_run("Figura: flujo Brain + MCP + agentes + revisión humana.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(88, 96, 105)

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_diagram()
    insert_into_docx()
    print(f"Updated {DOCX_PATH}")
    print(f"Diagram {DIAGRAM_PATH}")
