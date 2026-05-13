from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Brain - Producto y Arquitectura Base.docx"

ACCENT = "1F4E5F"
MUTED = "64748B"
LIGHT = "EAF2F4"
GRID = "CBD5E1"
TEXT = "111827"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = GRID, size: str = "4") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, *, bold=False, size=None, color=TEXT) -> None:
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    if size:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text: str = "", style: str | None = None, *, bold=False, color=TEXT, size=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        style_run(run, bold=bold, color=color, size=size)
    return p


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        run = p.add_run(item)
        style_run(run, size=10.5)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(item)
        style_run(run, size=10.5)


def add_callout(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_borders(cell, color="D7E2E8")
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    style_run(r, bold=True, size=11, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    style_run(r2, size=10.5, color=TEXT)
    doc.add_paragraph()


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, sum(widths))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        set_cell_width(hdr[idx], widths[idx])
        set_cell_shading(hdr[idx], LIGHT)
        set_cell_borders(hdr[idx])
        set_cell_margins(hdr[idx])
        hdr[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        style_run(run, bold=True, size=9.5, color=ACCENT)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_width(cells[idx], widths[idx])
            set_cell_borders(cells[idx])
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            style_run(run, size=9.2)
    doc.add_paragraph()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor.from_string(TEXT)

    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.05

    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string(ACCENT)
    styles["Title"].paragraph_format.space_after = Pt(8)

    styles["Subtitle"].font.size = Pt(11)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Subtitle"].paragraph_format.space_after = Pt(16)

    for h, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        styles[h].font.size = Pt(size)
        styles[h].font.bold = True
        styles[h].font.color.rgb = RGBColor.from_string(ACCENT)
        styles[h].paragraph_format.space_before = Pt(9 if h == "Heading 1" else 6)
        styles[h].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("Brain by DMKINGS - Producto y arquitectura base")
    style_run(run, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Documento interno + venta")
    style_run(run, size=8.5, color=MUTED)

    return doc


def build_doc() -> None:
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.add_run("Brain").bold = True
    title.runs[0].font.color.rgb = RGBColor.from_string(ACCENT)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Producto y arquitectura base para convertir una empresa en un sistema operativo AI-native")
    subtitle.runs[0].font.name = "Arial"
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    add_table(
        doc,
        ["Campo", "Detalle"],
        [
            ["Marca", "DMKINGS"],
            ["Uso", "Documento interno + venta"],
            ["Version", "v0.1 - 2026-05-13"],
            ["Base de negocio", "DMKINGS Blueprint Operativo y Pricing Model y Cobro"],
            ["Referencia tecnica", "Estructura repo-first apoyada en darLAAGAM/ai-native-playbook y usecompai.com"],
        ],
        [2100, 7260],
    )

    add_callout(
        doc,
        "Idea central",
        "Brain no es un chatbot ni un RAG. Es un repo operativo vivo que contiene la memoria, criterios, skills, prompts, plantillas, governance y acceso ordenado a las herramientas que ya usa la empresa.",
    )

    doc.add_heading("1. Que es Brain", level=1)
    add_paragraph(
        doc,
        "Brain es el producto de DMKINGS para empaquetar la forma de trabajar AI-native de una empresa. Convierte conocimiento disperso, procesos informales, criterios de decision, herramientas internas y aprendizajes repetibles en una base usable por LLMs y equipos reales.",
    )
    add_paragraph(
        doc,
        "La promesa no es vender tecnologia aislada. La promesa es que la empresa deje de depender de prompts sueltos, documentos perdidos o personas clave que guardan todo en la cabeza. Brain crea una memoria operativa versionada y accionable.",
    )
    add_bullets(
        doc,
        [
            "Funciona con Claude, Codex, ChatGPT, Cursor u otros clientes AI cuando el cliente ya tiene stack.",
            "Puede entregarse con OpenWork cuando el cliente necesita workspace propio, usuarios, permisos, historial y experiencia empaquetada.",
            "Agrupa las herramientas existentes de la empresa mediante una capa MCP para que el LLM/OpenWork pueda acceder a ellas de forma controlada.",
            "Deja preparado el terreno para vender despues RAGs, agentes recurrentes, automatizaciones o integraciones profundas como modulos separados.",
        ],
    )

    doc.add_heading("2. Posicionamiento de producto", level=1)
    add_callout(
        doc,
        "Mensaje comercial",
        "Brain convierte la empresa en un sistema operativo AI-native: memoria viva, skills propias y acceso ordenado a sus herramientas para que cualquier LLM pueda trabajar con contexto y accion real.",
    )
    add_table(
        doc,
        ["Brain es", "Brain no es"],
        [
            ["Un repo operativo vivo.", "Un chatbot generico."],
            ["Una biblioteca privada de skills, prompts y templates.", "Un RAG como producto principal."],
            ["Una forma de conectar herramientas existentes al LLM/OpenWork.", "Una automatizacion puntual."],
            ["Una base para que los proyectos a medida sean mas rapidos y reutilizables.", "Un pack de prompts sueltos."],
        ],
        [4680, 4680],
    )

    doc.add_heading("3. Repo y arquitectura Brain", level=1)
    add_paragraph(
        doc,
        "La estructura de Brain se apoya brevemente en el enfoque repo-first de darLAAGAM/ai-native-playbook y usecompai.com: playbook, skills, prompts, templates, memoria operativa y patrones versionados. DMKINGS lo adapta a una oferta vendible para pymes, especialmente empresas industriales y B2B.",
    )
    add_table(
        doc,
        ["Carpeta", "Funcion dentro de Brain"],
        [
            ["README.md", "Entrada del repo: que es Brain, como se usa y que incluye."],
            ["docs/", "Playbook operativo, principios, procesos, guias por rol y decisiones."],
            ["skills/", "Procedimientos ejecutables por LLM con objetivo, inputs, pasos, guardrails y output."],
            ["prompts/", "Master prompt, prompts por rol, prompts por equipo y custom instructions."],
            ["templates/", "Starters, onboarding, SOUL.md, checklists y plantillas de entrega."],
            ["memory/", "Decisiones, aprendizajes, gaps, changelog y world model de la empresa."],
            ["patterns/", "Patrones reutilizables de operacion y buenas practicas anonimizables."],
            ["governance/", "Permisos, privacidad, datos prohibidos, aprobaciones, limites y auditoria."],
            ["setup/", "Instalacion para Claude, Codex, ChatGPT, Cursor y OpenWork."],
            ["mcp/", "Configuracion y conectores que agrupan herramientas existentes para hacerlas accesibles al LLM/OpenWork."],
        ],
        [1850, 7510],
    )
    add_paragraph(
        doc,
        "La capa MCP no se vende como producto separado dentro de Brain. Es la forma practica de juntar las herramientas que el cliente ya usa y hacer que esten disponibles con permisos, limites y trazabilidad. La POC puede conectar 1-2 herramientas sencillas si procede; integraciones complejas se presupuestan aparte.",
    )

    doc.add_heading("4. Dos vias de venta", level=1)
    add_table(
        doc,
        ["Via", "Cuando usarla", "Entregable principal"],
        [
            [
                "Brain BYO-AI",
                "El cliente ya usa Claude, Codex, ChatGPT, Cursor u otra herramienta y quiere orden, contexto y rutinas propias.",
                "Repo Brain personalizado, master prompt, skills, templates, memoria, onboarding y setup para su stack actual.",
            ],
            [
                "Brain + OpenWork",
                "El cliente no tiene stack AI claro o necesita interfaz propia, usuarios, permisos, historial y experiencia empaquetada.",
                "OpenWork configurado con Brain, workspace, usuarios, permisos, historial, skills y acceso a herramientas conectadas.",
            ],
        ],
        [1900, 3650, 3810],
    )

    doc.add_heading("5. Delivery", level=1)
    add_table(
        doc,
        ["Fase", "Objetivo", "Entregable"],
        [
            ["Discovery", "Entender procesos, herramientas, conocimiento, riesgos, usuarios y quick wins.", "Mapa de Brain, scope de POC, datos necesarios y propuesta."],
            ["POC", "Crear una primera version usable de Brain con usuarios reales.", "Repo inicial, master prompt, 5-10 skills, memoria base, templates, setup y conexion minima si procede."],
            ["Handover", "Que el cliente lo use sin depender de DMKINGS para cada interaccion.", "Onboarding, guia de uso, limites, ejemplos y canal de soporte."],
            ["Retainer", "Mantener Brain vivo y convertir uso real en mejoras.", "Actualizacion mensual de skills, prompts, playbook, memoria y conexiones existentes."],
            ["Evolucion", "Detectar oportunidades adicionales sin mezclar el core con proyectos a medida.", "Backlog valorado de RAGs, agentes, automatizaciones, conectores o integraciones profundas."],
        ],
        [1550, 3900, 3910],
    )
    add_bullets(
        doc,
        [
            "Regla de POC: empezar con memoria, skills y setup; no con una integracion legacy compleja.",
            "Regla de alcance: Brain core ordena la empresa; los proyectos a medida se venden como modulos adicionales.",
            "Regla de seguridad: lectura antes que escritura y aprobacion humana para acciones sensibles.",
        ],
    )

    doc.add_heading("6. Paquetes y pricing", level=1)
    add_table(
        doc,
        ["Paquete", "Precio", "Incluye"],
        [
            ["Brain Discovery", "1.500-3.000 EUR", "Diagnostico, mapa de herramientas, procesos, conocimiento, riesgos, quick wins y propuesta de POC."],
            ["Brain POC", "4.000-12.000 EUR", "Repo Brain inicial, master prompt, 5-10 skills, memoria base, templates, onboarding, setup BYO-AI u OpenWork y conexion minima de 1-2 herramientas si procede."],
            ["Brain BYO-AI Retainer", "500-1.500 EUR/mes", "Mantenimiento de skills, prompts, playbook, memoria, mejoras mensuales, soporte ligero y ajustes sobre conexiones existentes."],
            ["Brain + OpenWork Setup", "5.000-10.000 EUR setup", "OpenWork empaquetado con marca, workspace, usuarios, permisos, historial, skills, configuracion y acceso a herramientas conectadas."],
            ["Brain + OpenWork Monthly", "750-1.500 EUR/mes", "Operacion, soporte, actualizaciones, control de versiones, onboarding y mantenimiento de experiencia."],
        ],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Extras fuera del core",
        "RAGs, agentes recurrentes, automatizaciones, conectores profundos e integraciones ERP/CRM complejas se presupuestan aparte. Brain prepara la base para venderlos mejor, pero no los absorbe dentro del precio core.",
    )

    doc.add_heading("7. Primer Brain Starter de DMKINGS", level=1)
    add_table(
        doc,
        ["Activo", "Estado minimo"],
        [
            ["Master prompt", "Prompt canonico con contexto, rol, principios, limites y forma de usar la memoria."],
            ["5-10 skills base", "Resumen documental, preparacion de reunion, email cliente, decision memo, checklist operativo, output record y health check."],
            ["Memoria base", "Decisiones, aprendizajes, gaps, world model y changelog."],
            ["Plantilla de onboarding", "Instrucciones para instalar y usar Brain en Claude/Codex/ChatGPT/OpenWork."],
            ["Governance basica", "Datos prohibidos, permisos, acciones sensibles y reglas de aprobacion."],
            ["Capa MCP minima", "Conexion sencilla de herramientas si procede, empezando por bajo riesgo y alto uso."],
            ["Variante OpenWork", "Workspace empaquetado con usuarios, permisos, historial y skills disponibles."],
        ],
        [2600, 6760],
    )

    doc.add_heading("8. Roadmap", level=1)
    add_table(
        doc,
        ["Periodo", "Objetivo", "Activos"],
        [
            ["Semana 1", "Cerrar repo starter y documento base.", "Estructura Brain, documento comercial, pricing y checklist de venta."],
            ["Semana 2", "Preparar demo y primeras skills.", "Master prompt, skills iniciales, onboarding y ejemplo de cliente ficticio."],
            ["30-60 dias", "Vender y ejecutar primera POC.", "Primer Brain real, feedback de usuarios, metricas de uso y mejoras."],
            ["3-6 meses", "Convertir aprendizaje en biblioteca reusable.", "Skills verticales, conectores reutilizables, templates de delivery y casos reales."],
        ],
        [1600, 3500, 4260],
    )

    doc.add_heading("9. Checklist de verificacion antes de vender", level=1)
    add_bullets(
        doc,
        [
            "El discurso deja claro que Brain no es RAG ni chatbot.",
            "La POC tiene un rango de 4.000-12.000 EUR y alcance cerrado.",
            "OpenWork aparece como segunda via de entrega, no como requisito obligatorio.",
            "La capa MCP se explica como acceso ordenado a herramientas existentes, no como producto separado.",
            "No se promete escritura en sistemas sensibles sin aprobaciones y logs.",
            "RAGs, agentes recurrentes, automatizaciones e integraciones complejas aparecen como upsells fuera del core.",
            "El repo Brain futuro queda claro y versionable.",
        ],
    )

    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
