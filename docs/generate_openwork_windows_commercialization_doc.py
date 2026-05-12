from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).with_name("openwork_windows_selfhosted_comercializacion.docx")


ACCENT = RGBColor(32, 86, 142)
MUTED = RGBColor(95, 105, 117)
LIGHT = "EEF3F8"
SOFT = "F7F9FB"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="D6DEE8", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    widths_dxa = [int(round(width * 1440)) for width in widths]
    total_dxa = sum(widths_dxa)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def style_run(run, bold=False, color=None, size=None):
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.color.rgb = ACCENT if level == 1 else RGBColor(25, 33, 42)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.add_run(item)


def add_callout(doc, title, body, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    set_table_borders(table, color="C9D6E4")
    set_table_geometry(table, [6.5])
    p = cell.paragraphs[0]
    r = p.add_run(title)
    style_run(r, bold=True, color=ACCENT, size=11)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        style_run(r, bold=True, color=RGBColor(24, 40, 56), size=10)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(text)
    set_table_borders(table)
    if widths:
        set_table_geometry(table, widths)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Title", 22, ACCENT),
        ("Subtitle", 11, MUTED),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 12.5, RGBColor(25, 33, 42)),
        ("Heading 3", 11, RGBColor(25, 33, 42)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(4)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "OpenWork - Windows, self-hosted y comercializacion"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.runs[0].font.name = "Arial"
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = MUTED

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run("Documento de trabajo - ").font.size = Pt(8)
    fp.add_run(str(date.today())).font.size = Pt(8)

    title = doc.add_paragraph(style="Title")
    title.add_run("Matrix/OpenWork: producto empaquetado, Windows y comercializacion")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Documento vivo para ordenar la idea: Matrix como build propia paquetizada, skills premium por sector, incompatibilidades Windows, servidores, modelos y vias de salida al mercado."
    )

    meta = add_table(
        doc,
        ["Decision", "Lectura actual"],
        [
            ["Producto", "Matrix como version propia basada en OpenWork, empaquetada para Windows y personalizada con skills/workflows de alto valor."],
            ["Tension principal", "Desktop local aporta confianza y control; self-hosted/cloud aportan acceso multiplataforma y actualizaciones mas rapidas."],
            ["Riesgo Windows", "Muchos clientes usan Windows, pero parte del ecosistema agentico y dev tooling nace primero en Linux/macOS."],
            ["Hipotesis comercial", "No vender una app generica de agentes: vender paquetes sectoriales listos para operar, con instalacion, skills, soporte e integraciones."],
        ],
        widths=[1.75, 4.75],
    )

    add_heading(doc, "1. Introduccion: Matrix como producto paquetizado")
    doc.add_paragraph(
        "La idea comercial no es limitarse a revender una build de OpenWork ni entregar un repositorio tecnico. Matrix seria una version propia basada en el codigo open source de OpenWork, compilada, empaquetada y adaptada para clientes concretos. El valor no estaria solo en la app, sino en convertir una tecnologia agentica general en una solucion operativa que el cliente pueda instalar, abrir y usar."
    )
    doc.add_paragraph(
        "La capa diferencial son las skills super top: instrucciones, criterios, pasos, validaciones, formatos de salida, ejemplos y workflows disenados para un sector o para un cliente. En vez de vender 'un chat con agentes', Matrix puede vender 'un sistema que genera descripciones de inmuebles', 'resume expedientes', 'prepara respuestas a leads', 'audita carpetas de documentos' o 'convierte PDFs en tareas y emails listos para revisar'."
    )
    add_table(
        doc,
        ["Elemento vendible", "Que recibe el cliente", "Por que tiene valor"],
        [
            ["Instalador Windows", "Una app con nombre, icono, configuracion y binarios incluidos.", "Reduce friccion: el cliente no instala Node, pnpm, Bun, OpenCode ni herramientas sueltas."],
            ["Skills premium", "Paquetes de criterio operativo por sector o por proceso.", "Convierte la IA generica en resultados repetibles y alineados al negocio."],
            ["Workflows", "Flujos ejecutables con pasos, aprobaciones, plantillas y formatos.", "Pasa de conversacion suelta a produccion diaria."],
            ["Integraciones", "Conexiones con Drive, email, CRM, hojas, APIs o mensajeria.", "Matrix entra en el sistema real de trabajo del cliente."],
            ["Soporte y mejora", "Instalacion, formacion, actualizaciones y ajuste mensual.", "Hace viable venderlo como servicio recurrente, no solo como entrega unica."],
        ],
        widths=[1.5, 2.55, 2.45],
    )
    add_callout(
        doc,
        "Tesis comercial",
        "El cliente no compra OpenWork ni una herramienta de agentes. Compra una solucion Matrix adaptada a su negocio, empaquetada para su entorno y mantenida para que produzca resultados sin obligarle a entender el stack tecnico.",
    )

    add_heading(doc, "2. Personalizacion como ventaja comercial")
    doc.add_paragraph(
        "Si Matrix se quiere vender, la personalizacion debe ir mas alla de cambiar el logo. La promesa comercial fuerte es que cada cliente sienta que Matrix entiende sus herramientas, reglas, tono, procesos y marca. El objetivo no es ofrecer un agente generico, sino un sistema operativo personalizado para trabajar con agentes dentro de un negocio concreto."
    )
    add_table(
        doc,
        ["Capa", "Ejemplos", "Valor comercial"],
        [
            ["Visual", "Nombre, logo, icono, colores, tipografia, tema claro/oscuro, pantalla inicial y dashboard por rol.", "Hace que la herramienta se perciba como propia, especialmente en marca blanca."],
            ["Agente", "Perfiles como CTO, product manager, soporte, ventas, analista o asistente documental.", "El usuario no elige un modelo; elige un rol operativo que entiende su trabajo."],
            ["Tono y estilo", "Directo, pedagogico, formal, creativo, tecnico, ejecutivo o adaptado al manual de marca.", "Alinea respuestas, documentos y mensajes con la comunicacion real del cliente."],
            ["Autonomia", "Solo sugerir, pedir permiso, ejecutar cambios pequenos o trabajar en modo mas automatico.", "Permite vender confianza: cada cliente decide cuanto control cede al agente."],
            ["Workspace", "AGENTS.md, skills, permisos, prompts, proveedores y modelos permitidos por cliente.", "Convierte cada workspace en una configuracion operativa aislada y vendible."],
            ["Conocimiento", "Documentacion interna, FAQs, procesos, plantillas, historico de casos y ejemplos reales.", "Matrix aprende la forma de trabajar del cliente y reduce la friccion de adopcion."],
        ],
        widths=[1.25, 3.05, 2.2],
    )
    add_callout(
        doc,
        "Frase de venta",
        "Matrix se adapta a las herramientas, reglas, tono, workflows y marca de tu equipo.",
        fill=SOFT,
    )

    add_heading(doc, "3. Ediciones verticales y workflows propios")
    doc.add_paragraph(
        "Una estrategia potente es crear ediciones verticales. Todas comparten la misma base tecnica, pero cambian skills, workflows, conectores, pantallas iniciales, permisos y plantillas. Asi Matrix se puede vender por sector, no como una app horizontal dificil de explicar."
    )
    add_table(
        doc,
        ["Edicion", "Casos de uso principales"],
        [
            ["Agencias", "Briefs, landings, propuestas, reporting, campanas y contenido multicanal."],
            ["Inmobiliarias", "Anuncios, respuestas a leads, comparativas, Idealista, CRM y documentacion comercial."],
            ["SaaS", "Issues, PRs, soporte, changelogs, onboarding, releases y documentacion tecnica."],
            ["Ecommerce", "Fichas de producto, atencion al cliente, inventario, campanas y analisis de pedidos."],
            ["Despachos", "Borradores, revision documental, expedientes, compliance y comunicaciones."],
            ["Educacion", "Tutoria, materiales, seguimiento de alumnos, evaluaciones y comunicacion."],
        ],
        widths=[1.55, 4.95],
    )
    doc.add_paragraph(
        "El mayor valor percibido aparece cuando Matrix ofrece botones y acciones que el usuario reconoce como parte de su dia a dia: crear PR, responder cliente, generar informe, revisar contrato, preparar release, importar inmueble, clasificar lead, resumir expediente, preparar propuesta o generar checklist."
    )
    add_bullets(
        doc,
        [
            "Cada boton puede ejecutar un workflow completo: leer datos, razonar con contexto, usar skills concretas, pedir aprobacion humana y crear un documento, tarea, email o PR.",
            "Los conectores elevan el valor: GitHub, GitLab, Linear, Jira, Slack, Teams, WhatsApp, Telegram, Drive, OneDrive, Notion, Airtable, Sheets, Gmail, Outlook, CRM, Shopify, WooCommerce, Vercel, Supabase o AWS.",
            "La personalizacion de permisos es diferencial en empresas: roles, aprobaciones, sandbox, auditoria y politicas como no tocar produccion o no enviar emails sin revision.",
        ],
    )
    add_table(
        doc,
        ["Paquete", "Incluye"],
        [
            ["Starter", "App local, agentes basicos, personalizacion visual ligera y skills iniciales."],
            ["Pro", "Skills personalizadas, workflows, memoria por workspace y conectores basicos."],
            ["Team", "Roles, permisos, colaboracion, auditoria y plantillas compartidas."],
            ["Enterprise", "Marca blanca, self-hosting, SSO, politicas, conectores privados y soporte premium."],
            ["Vertical editions", "Paquetes especificos como Matrix para inmobiliarias, agencias, SaaS, ecommerce o despachos profesionales."],
        ],
        widths=[1.65, 4.85],
    )

    add_heading(doc, "4. Skills personalizables: producto vendible y no solo prompts")
    doc.add_paragraph(
        "Las skills personalizables son una de las piezas mas vendibles de Matrix/OpenWork porque convierten conocimiento operativo en capacidad reutilizable. Una skill no deberia entenderse como un prompt suelto, sino como un pequeno manual de ejecucion: cuando usarla, que fuentes mirar, que pasos seguir, que validaciones aplicar, que formato entregar y donde pedir aprobacion humana."
    )
    doc.add_paragraph(
        "Esto permite vender paquetes por sector o cliente sin rehacer todo el producto. La misma base puede cargar una carpeta `.opencode/skills/`, comandos rapidos, agentes por rol y configuracion de workspace para que el sistema arranque ya con criterio de negocio."
    )
    add_table(
        doc,
        ["Componente", "Que contiene", "Valor para el cliente"],
        [
            ["Skill", "Instrucciones, criterios, fuentes, pasos, guardrails y formato de salida.", "Hace repetible una tarea compleja y reduce dependencia de que el usuario sepa pedir bien."],
            ["Command", "Atajo tipo `/ordenar-expediente` o `/preparar-paquete-cliente` que invoca una rutina.", "Convierte workflows frecuentes en acciones visibles y faciles de enseñar."],
            ["Agente", "Rol operativo con tono, permisos, prudencia y ambito definidos.", "El usuario elige una responsabilidad, no un modelo ni una configuracion tecnica."],
            ["Workspace", "Skills, commands, agentes, MCP, proveedores, reglas y memoria aisladas por cliente.", "Cada cliente recibe una version portable, auditable y ajustable de Matrix."],
            ["Hot reload", "Cambios en `.opencode/` que pueden recargarse sin reconstruir la app completa.", "Permite iterar skills durante pilotos y convertir aprendizaje real en mejoras rapidas."],
        ],
        widths=[1.35, 2.7, 2.45],
    )
    add_bullets(
        doc,
        [
            "Las skills pueden empaquetar conocimiento experto: taxonomias, checklists, tono de comunicacion, reglas de compliance, ejemplos buenos y anti-ejemplos.",
            "Los commands hacen que esas capacidades aparezcan como botones mentales: ordenar, clasificar, auditar, sincronizar, preparar, revisar o publicar.",
            "La venta mejora cuando se entrega una biblioteca inicial y un proceso mensual de refinamiento: observar tareas reales, crear skills nuevas, medir calidad y bloquear acciones sensibles hasta aprobacion.",
        ],
    )
    add_callout(
        doc,
        "Idea clave",
        "La skill personalizada es el envase del know-how. Matrix vende criterio operativo empaquetado, no solo acceso a modelos.",
        fill=SOFT,
    )

    add_heading(doc, "5. Ejemplo real: Alcuaz Temino AI")
    doc.add_paragraph(
        "En `C:\\matrix-alcuaz-temino` se preparo una demo white-label llamada Alcuaz Temino AI. La idea no era vender un chatbot juridico, sino un operador documental para despacho: un sistema que ayuda a convertir carpetas, emails, documentos, sentencias y comunicaciones en expedientes ordenados, trazables y listos para revision profesional."
    )
    add_table(
        doc,
        ["Area", "Personalizacion realizada"],
        [
            ["Marca y experiencia", "Nombre visible Alcuaz Temino AI, estetica sobria/premium, logo y wordmark del despacho, starters orientados a workflows y limpieza de ruido secundario para demo."],
            ["Agente", "Agente `abogado-operativo`, orientado a ejecucion prudente: fuente, destino, permisos, criterio, acciones propuestas, riesgos y revision humana."],
            ["Skills", "`ordenar-expediente`, `clasificar-sentencias`, `timeline-probatorio`, `paquete-cliente`, `sincronizar-drive-sharepoint`, `auditoria-documental-rgpd` y `triage-correo-despacho`."],
            ["Commands", "`/ordenar-expediente`, `/clasificar-sentencias`, `/crear-timeline`, `/preparar-paquete-cliente`, `/sincronizar-documentos` y `/auditar-rgpd-expediente`."],
            ["Conectores recomendados", "Drive o SharePoint, Gmail u Outlook, calendario, WhatsApp/Telegram via router, Teams/Slack, OCR/document AI, CRM ligero y base privada de conocimiento."],
        ],
        widths=[1.55, 4.95],
    )
    add_table(
        doc,
        ["Workflow demo", "Salida esperada sin credenciales reales"],
        [
            ["Ordenar expediente", "Inventario documental, taxonomia, nombres, carpetas destino y tabla origen -> destino."],
            ["Clasificar sentencias", "Jurisdiccion, organo, fecha, fallo, ratio decidendi, cuantias, utilidad y confianza."],
            ["Timeline probatorio", "Fecha, hecho, documento soporte, fuerza probatoria, laguna y accion siguiente."],
            ["Paquete cliente", "Resumen ejecutivo, checklist de documentos pendientes, email y WhatsApp para aprobacion."],
            ["Sincronizacion documental", "Estructura destino, permisos, conflictos, duplicados y pasos de ejecucion para Drive/SharePoint."],
            ["Auditoria RGPD", "Categorias de datos, datos sensibles, finalidad, accesos, retencion y medidas."],
        ],
        widths=[1.75, 4.75],
    )
    add_callout(
        doc,
        "Leccion comercial de Alcuaz Temino",
        "Incluso sin credenciales reales, la demo ya puede enseñar planes ejecutables y trazables. El sistema no finge haber movido documentos: propone la accion, muestra riesgos y pide aprobacion.",
    )
    add_bullets(
        doc,
        [
            "Fase demo: branding, starters, skills y commands documentales, documento comercial y modo plan ejecutable.",
            "Fase piloto: conectar Drive o SharePoint, probar expedientes anonimizados, ajustar taxonomia real y medir tiempo ahorrado.",
            "Fase premium: email, calendario, WhatsApp, OCR, timeline automatico, portal cliente, auditoria, permisos por rol y reporting de productividad.",
        ],
    )

    add_heading(doc, "6. Posibles mejoras sobre el OpenCode actual")
    doc.add_paragraph(
        "Ademas de la personalizacion visual y de las skills personalizadas, Matrix puede diferenciarse tocando codigo del producto para mejorar la experiencia base de OpenCode/OpenWork. La idea no es convertir Matrix en un CRM, ClickUp o software vertical clasico, sino en una capa mas potente para manejar el PC, el navegador, las tools y el negocio con lenguaje natural."
    )
    add_table(
        doc,
        ["Mejora", "Que aportaria"],
        [
            ["UI en skills", "Una skill podria exponer inputs, formularios, checklist, progreso, preview de resultado y botones de aprobacion, ademas de seguir disponible por chat o voz."],
            ["Integracion profunda con navegador", "Controlar sesiones, observar tabs, extraer datos, rellenar formularios, descargar archivos, recordar rutinas web y repetirlas con aprobacion."],
            ["Memoria como autoaprendizaje", "Capturar resultados, pasos, tools usadas, errores, permisos y feedback para proponer memorias, reglas, skills o automations nuevas."],
            ["Autoaprendizaje gobernado", "El agente observa y propone, pero pide aprobacion antes de guardar memoria, crear/refinar skills o activar automations."],
            ["Skill proposals", "Al terminar una sesion, sugerir guardar la rutina como skill, crear una accion rapida o recordar una regla."],
            ["Skill diff", "Mostrar que cambio en una skill, por que, de que sesion viene y permitir aceptar, editar o rechazar."],
            ["Memory inbox", "Bandeja de aprendizajes pendientes: preferencias, rutinas, errores recurrentes, pasos de navegador y permisos habituales."],
            ["Run recorder", "Registro estructurado de cada ejecucion: objetivo, pasos, tools, archivos/webs tocadas, aprobaciones, resultado y evaluacion."],
            ["Pattern detector", "Detectar tareas repetidas, formatos corregidos varias veces, portales usados de forma recurrente o skills que fallan por el mismo motivo."],
            ["Tool readiness", "Cuando Matrix no pueda hacer algo, proponer que falta: MCP, tool local, skill, permiso, navegador o credencial."],
            ["Browser memory", "Recordar como se opera una web: donde estan los leads, que boton exporta, que filtro usar o donde se descarga un informe."],
            ["Automations desde memoria", "Convertir rutinas repetidas en tareas programadas o disparadas por evento, manteniendo aprobaciones para acciones sensibles."],
        ],
        widths=[1.9, 4.6],
    )
    add_callout(
        doc,
        "Propuesta fuerte",
        "Matrix no reemplaza tus herramientas. Aprende como usas tu PC, navegador y tools, y convierte trabajo repetido en capacidades reutilizables.",
        fill=SOFT,
    )

    add_heading(doc, "7. Resumen ejecutivo")
    doc.add_paragraph(
        "La idea fuerte es que OpenWork no debe depender de una sola forma de ejecucion. El cliente puede querer una app local, un servidor propio o una experiencia cloud. Windows es central porque una gran parte del mercado potencial trabaja en Windows, pero tambien es el entorno donde mas se notan las incompatibilidades de herramientas CLI, permisos, shell, rutas, dependencias nativas, Docker y modelos locales."
    )
    add_callout(
        doc,
        "Tesis",
        "Empaquetar para reducir friccion, pero desacoplar runtimes y workers para no congelar a los clientes en una version peor cuando el proyecto upstream mejore.",
    )

    add_heading(doc, "8. Punto de partida: que problema estamos resolviendo")
    add_bullets(
        doc,
        [
            "Cowork/Codex-like workflows son potentes, pero muchas soluciones son cerradas, caras o atadas a un proveedor concreto.",
            "Usuarios no tecnicos quieren ejecutar trabajo agentico sin pelearse con terminales, tokens, rutas y entornos.",
            "Equipos tecnicos quieren control: self-hosting, logs, permisos, workers, modelos configurables y portabilidad.",
            "La experiencia debe sentirse premium, no como un panel utilitario de infraestructura.",
        ],
    )

    add_heading(doc, "9. Casuistica Windows: desde la incompatibilidad hasta la adopcion")
    doc.add_paragraph(
        "Windows no es un detalle secundario: es una condicion comercial. Si OpenWork aspira a usuarios de negocio, founders, equipos pequenos y empresas tradicionales, Windows debe estar contemplado desde el principio."
    )
    add_table(
        doc,
        ["Caso Windows", "Riesgo", "Via de solucion"],
        [
            ["CLI y shell", "Comandos pensados para bash fallan en PowerShell/cmd; quoting y rutas cambian.", "Abstraer comandos frecuentes, perfiles por shell y guias especificas para PowerShell."],
            ["Rutas y permisos", "Separadores, espacios, UAC, antivirus y directorios protegidos rompen flujos.", "Usar rutas de datos propias, checks previos y mensajes claros de permisos."],
            ["Dependencias nativas", "Node, Python, Git, compilers y binarios pueden no estar instalados o variar.", "Empaquetar runtimes minimos o usar sidecars versionados."],
            ["Docker", "Docker Desktop requiere WSL2 y permisos; en empresas puede estar bloqueado.", "Ofrecer Docker como camino self-hosted, no como unica forma local."],
            ["Modelos locales", "VRAM/RAM insuficiente, drivers distintos y rendimiento irregular.", "Modelo hibrido: gratis/rapido por defecto y API/modelos top bajo demanda."],
            ["Actualizaciones", "Instaladores y auto-update pueden chocar con politicas IT.", "Canales estable/beta, actualizador firmado y releases offline para empresas."],
        ],
        widths=[1.55, 2.35, 2.6],
    )

    add_heading(doc, "10. Modelos locales, gratis y modelos top via API")
    doc.add_paragraph(
        "Los modelos que aparecen como MiniMax, DeepSeek Flash, Ring, Nemotron o similares probablemente vienen de un proveedor/router, no necesariamente de ejecucion local pura. Para producto, conviene separar tres conceptos: modelo local real, modelo remoto gratuito y modelo top via API."
    )
    add_table(
        doc,
        ["Tipo", "Cuándo usarlo", "Trade-off"],
        [
            ["Gratis/rapido", "Chat, resumen, navegacion de codigo y tareas simples.", "Menor coste, pero mas alucinaciones y peor razonamiento profundo."],
            ["Top via API", "Arquitectura, debugging dificil, PR review, refactors y decisiones de alto coste.", "Calidad alta, pero coste variable y dependencia externa."],
            ["Local puro", "Privacidad, demos offline, entornos bloqueados o coste fijo.", "Requiere hardware, configuracion y no siempre alcanza calidad cloud."],
        ],
        widths=[1.45, 3.05, 2.0],
    )
    add_callout(
        doc,
        "Estrategia recomendada",
        "Default barato o gratis para tareas corrientes; escalada manual a modelos top cuando el coste de equivocarse sea mayor que el coste de la API.",
        fill=SOFT,
    )

    add_heading(doc, "11. Empaquetado: el problema de quedarse atras")
    doc.add_paragraph(
        "La pega del empaquetado es que una version incluida dentro de la app queda congelada hasta la siguiente release. Si OpenCode, un provider, un modelo o una herramienta mejora, el cliente no recibe esa mejora automaticamente. El empaquetado da confianza, pero puede crear clientes con una experiencia peor que upstream."
    )
    add_table(
        doc,
        ["Estrategia", "Ventaja", "Riesgo"],
        [
            ["Runtime empaquetado", "Funciona out-of-the-box, instalacion sencilla.", "Se queda viejo si upstream avanza rapido."],
            ["Runtime actualizable", "Mejoras sin reinstalar toda la app.", "Requiere versionado, rollback y validacion."],
            ["Runtime externo compatible", "Power users usan su propio OpenCode/worker.", "Mas superficie de soporte y configuracion."],
            ["Hosted worker", "Control de version, performance y soporte centralizados.", "Mas responsabilidad en seguridad, coste e infraestructura."],
        ],
        widths=[1.85, 2.35, 2.3],
    )
    add_bullets(
        doc,
        [
            "Regla de producto: empaquetar para fiabilidad, desacoplar para frescura.",
            "La UI debe hablar con contratos de API/version, no con detalles internos fragiles.",
            "Debe existir fallback al runtime conocido-bueno si una actualizacion falla.",
        ],
    )

    add_heading(doc, "12. Self-hosted: que significa realmente")
    doc.add_paragraph(
        "Self-hosted significa que el cliente instala OpenWork/OpenCode en su propio servidor y accede desde navegador. Esto resuelve una parte importante del problema Windows: el usuario puede estar en Windows, Mac, Linux o movil, porque el trabajo pesado ocurre en el servidor."
    )
    add_table(
        doc,
        ["Modelo", "Para quien", "Implicacion"],
        [
            ["Desktop local", "Usuarios individuales y equipos pequenos que trabajan con archivos locales.", "Maxima sensacion local-first; requiere builds por sistema operativo."],
            ["Self-hosted", "Empresas o equipos con IT que quieren control.", "Acceso por navegador; el servidor toca repos, secretos y comandos."],
            ["Hosted cloud", "Usuarios que quieren cero instalacion.", "Mayor comodidad; requiere billing, seguridad, soporte y escalado."],
        ],
        widths=[1.55, 2.55, 2.4],
    )

    add_heading(doc, "13. Docker Compose: como se instalaria")
    doc.add_paragraph(
        "Para self-hosted, Docker Compose es el camino mas universal. No es doble click, pero es familiar para equipos tecnicos y portable entre VPS, servidores Linux, Proxmox, NAS o entornos corporativos compatibles."
    )
    add_numbered(
        doc,
        [
            "El cliente instala Docker y Docker Compose en el servidor.",
            "Recibe una carpeta con docker-compose.yml, .env.example y README.md.",
            "Copia .env.example a .env y configura puerto, dominio, claves API, almacenamiento y politica de workers.",
            "Ejecuta docker compose up -d.",
            "Abre la URL interna o publica: http://servidor:puerto o https://openwork.empresa.com.",
            "Configura backups, reverse proxy/HTTPS y rotacion de tokens segun el entorno.",
        ],
    )
    add_callout(
        doc,
        "Para empresas",
        "Ofrecer siempre instalacion manual documentada ademas de un script tipo curl | sh. Muchas organizaciones no aceptan scripts remotos por politica de seguridad.",
    )

    add_heading(doc, "14. Seguridad y confianza operacional")
    doc.add_paragraph(
        "El punto delicado del self-hosted/cloud es que el servidor puede acceder a repos, tokens, archivos y comandos. La propuesta comercial no puede ser solo comodidad: debe incluir control, trazabilidad y limites."
    )
    add_bullets(
        doc,
        [
            "Permisos explicitos por workspace, repo y accion.",
            "Aprobaciones visibles para comandos peligrosos o acceso a secretos.",
            "Logs auditables de sesiones, herramientas, workers y cambios.",
            "Aislamiento por workspace y, si hay cloud, por tenant.",
            "Tokens revocables y separacion entre identidad de usuario, worker y proveedor de modelo.",
            "Modo empresa: actualizaciones controladas, pinning de versiones y despliegue offline si hace falta.",
        ],
    )

    add_heading(doc, "15. Posibles vias de comercializacion")
    add_table(
        doc,
        ["Via", "Oferta", "Cliente objetivo", "Notas"],
        [
            ["Desktop freemium", "App local gratis con modelos/proveedores configurables.", "Usuarios individuales, makers, devs.", "Buena entrada de mercado; monetizar por cloud, sync o features pro."],
            ["Pro personal", "App desktop + perfiles de modelos + automatizaciones + soporte prioritario.", "Power users y freelancers.", "Precio fijo mensual bajo; evita depender solo de margen API."],
            ["Self-hosted licencia", "Docker Compose/binario + actualizaciones + soporte.", "Equipos tecnicos y pymes.", "Cobro anual por instancia, usuarios o soporte."],
            ["Enterprise self-hosted", "SSO, auditoria, politicas, despliegue offline y SLA.", "Empresas con compliance.", "Mayor ticket; requiere documentacion y seguridad maduras."],
            ["Hosted cloud", "OpenWork gestionado con workers remotos.", "Usuarios que no quieren instalar nada.", "Mayor comodidad y mejor control de version; coste infra y seguridad mas altos."],
            ["Marketplace/partners", "Plantillas, skills, conectores y workers certificados.", "Ecosistema y consultoras.", "Puede convertirse en moat si hay buena distribucion."],
        ],
        widths=[1.35, 1.7, 1.65, 1.8],
    )

    add_heading(doc, "16. Recomendacion de posicionamiento inicial")
    add_bullets(
        doc,
        [
            "No vender Matrix/OpenWork como un simple wrapper de OpenCode; venderlo como una solucion empaquetada con criterio, skills, seguridad y experiencia.",
            "Priorizar Windows como argumento comercial: una experiencia agentica que no obliga al usuario a vivir en Linux/macOS.",
            "Ofrecer tres caminos claros: Desktop, Self-hosted y Hosted Cloud.",
            "Mantener un runtime empaquetado estable, pero permitir workers/runtimes actualizables y compatibles.",
            "Diseñar la narrativa comercial alrededor de confianza: local-first cuando importa, cloud-ready cuando conviene.",
        ],
    )

    add_heading(doc, "17. Roadmap sugerido")
    add_table(
        doc,
        ["Fase", "Objetivo", "Resultado esperado"],
        [
            ["0", "Desktop Windows usable con runtime conocido-bueno.", "Demo local fiable para usuarios no tecnicos."],
            ["1", "Perfiles de modelo y escalada manual a modelos top.", "Coste controlado y mejor calidad cuando importa."],
            ["2", "Docker Compose self-hosted con README solido.", "Primer despliegue en servidor de cliente tecnico."],
            ["3", "Actualizador de runtimes/workers con rollback.", "Clientes no quedan atados a versiones viejas."],
            ["4", "Cloud hosted y workers remotos.", "Comercializacion SaaS y acceso movil/navegador."],
            ["5", "Enterprise: SSO, audit logs, politicas y soporte.", "Venta B2B con confianza operacional."],
        ],
        widths=[0.7, 3.0, 2.8],
    )

    add_heading(doc, "18. Preguntas abiertas")
    add_bullets(
        doc,
        [
            "Cual es el primer segmento de pago: usuarios individuales, pymes tecnicas o empresas con IT?",
            "Que parte debe ser open-source y que parte puede ser cloud/pro/enterprise?",
            "Cual es el contrato minimo entre UI, servidor y worker para que no haya lock-in interno?",
            "Hasta que punto se quiere soportar local puro en Windows frente a empujar self-hosted/cloud?",
            "Que proveedor/modelo se recomienda por defecto sin comprometer neutralidad?",
        ],
    )

    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
